"""
AI Resume Builder - Main Application
Flask + MongoDB Atlas + Groq AI + Google OAuth
"""

import os
import json
import re
import shutil
import tempfile
import bcrypt
import requests
from datetime import datetime, timedelta
from functools import wraps
from bson import ObjectId
from bson.json_util import dumps as bson_dumps

from flask import (
    Flask, render_template, request, redirect, url_for,
    session, jsonify, flash, g
)
from flask.json.provider import DefaultJSONProvider
from flask_session import Session
from dotenv import load_dotenv
from pymongo import MongoClient
from groq import Groq
from authlib.integrations.requests_client import OAuth2Session

# ── PDF / DOCX parsing ──────────────────────────────────────────────────────
try:
    import fitz  # PyMuPDF
    PYMUPDF_AVAILABLE = True
except ImportError:
    PYMUPDF_AVAILABLE = False

try:
    import PyPDF2
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from pdfminer.high_level import extract_text as pdfminer_extract
    PDFMINER_AVAILABLE = True
except ImportError:
    PDFMINER_AVAILABLE = False

try:
    import pytesseract
    from PIL import Image
    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False

# ── Load environment ─────────────────────────────────────────────────────────
load_dotenv()


# ── Custom JSON Provider (Flask 3.x compatible) ──────────────────────────────
class MongoJSONProvider(DefaultJSONProvider):
    def default(self, obj):
        if isinstance(obj, ObjectId):
            return str(obj)
        if isinstance(obj, datetime):
            return obj.isoformat()
        return super().default(obj)


# ── Flask app setup ──────────────────────────────────────────────────────────
app = Flask(__name__)
app.json_provider_class = MongoJSONProvider
app.json = MongoJSONProvider(app)

app.secret_key = os.getenv("SECRET_KEY", "dev-secret-key-change-in-production")
app.config["SESSION_TYPE"] = "filesystem"
app.config["SESSION_PERMANENT"] = True
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=7)
app.config["SESSION_FILE_DIR"] = os.path.join(tempfile.gettempdir(), "flask_sessions")
os.makedirs(app.config["SESSION_FILE_DIR"], exist_ok=True)
Session(app)

# ── MongoDB ──────────────────────────────────────────────────────────────────
mongo_uri = os.getenv("MONGO_URI")
mongo_client = MongoClient(mongo_uri)
db = mongo_client["ai_resume_builder"]
users_col = db["users"]
resumes_col = db["resumes"]

# Indexes
users_col.create_index("email", unique=True)
resumes_col.create_index("user_id")

# ── Groq client ──────────────────────────────────────────────────────────────
groq_client = Groq(api_key=os.getenv("GROQ_API_KEY", ""))

# ── Google OAuth config ──────────────────────────────────────────────────────
GOOGLE_CLIENT_ID     = os.getenv("GOOGLE_CLIENT_ID", "")
GOOGLE_CLIENT_SECRET = os.getenv("GOOGLE_CLIENT_SECRET", "")
GOOGLE_REDIRECT_URI  = os.getenv("GOOGLE_REDIRECT_URI")
GOOGLE_AUTH_URL      = "https://accounts.google.com/o/oauth2/v2/auth"
GOOGLE_TOKEN_URL     = "https://oauth2.googleapis.com/token"
GOOGLE_USERINFO_URL  = "https://www.googleapis.com/oauth2/v2/userinfo"


# ── Helper: safe groq call ───────────────────────────────────────────────────
def groq_chat(messages, temperature=0.7, max_tokens=2048, response_format=None):
    """Call Groq API safely, optionally enforcing a structured response format."""
    try:
        kwargs = {
            "model": os.getenv("GROQ_MODEL", "openai/gpt-oss-120b"),
            "messages": messages,
            "temperature": temperature,
            "max_tokens": max_tokens,
        }
        if response_format is not None:
            kwargs["response_format"] = response_format

        resp = groq_client.chat.completions.create(**kwargs)
        content = resp.choices[0].message.content or ""
        return content.strip()
    except Exception as e:
        app.logger.error(f"Groq error: {e}")
        return ""


# ── Auth decorator ───────────────────────────────────────────────────────────
def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if "user_id" not in session:
            if request.is_json:
                return jsonify({"error": "Authentication required"}), 401
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated


def get_current_user():
    if "user_id" not in session:
        return None
    try:
        return users_col.find_one({"_id": ObjectId(session["user_id"])})
    except Exception:
        return None


# ── ATS domain keywords ───────────────────────────────────────────────────────
ATS_DOMAINS = {
    "Software Engineering": ["python","java","javascript","typescript","react","node","sql","docker","kubernetes","git","api","rest","microservices","cloud","aws","gcp","azure","ci/cd","agile","scrum"],
    "Data Science": ["python","r","machine learning","deep learning","tensorflow","pytorch","pandas","numpy","scikit-learn","statistics","sql","tableau","power bi","data visualization","nlp","computer vision"],
    "Product Management": ["roadmap","stakeholder","agile","scrum","kpis","okrs","user stories","sprint","backlog","mvp","a/b testing","metrics","jira","confluence","product strategy"],
    "Marketing": ["seo","sem","ppc","content marketing","social media","email marketing","google analytics","crm","hubspot","salesforce","branding","copywriting","campaigns","roi","lead generation"],
    "Finance": ["financial modeling","excel","bloomberg","valuation","dcf","financial statements","accounting","gaap","ifrs","risk management","portfolio","investment","cfa","cpa","audit"],
    "Healthcare": ["ehr","emr","clinical","hipaa","patient care","medical terminology","healthcare management","nursing","pharmacy","diagnosis","treatment","epic","cerner","billing","coding"],
    "Design": ["figma","sketch","adobe xd","ui/ux","wireframe","prototype","user research","design system","typography","color theory","photoshop","illustrator","indesign","branding","accessibility"],
    "DevOps": ["docker","kubernetes","jenkins","terraform","ansible","ci/cd","linux","bash","monitoring","prometheus","grafana","elk","nginx","aws","azure","gcp","helm","git","infrastructure"],
}


def compute_ats_score(resume_data, domain="Software Engineering"):
    """Compute ATS score based on keyword matching."""
    keywords = ATS_DOMAINS.get(domain, ATS_DOMAINS["Software Engineering"])
    text = " ".join([
        resume_data.get("summary", ""),
        " ".join(resume_data.get("skills", [])),
        " ".join([
            exp.get("description", "") + " " + exp.get("title", "")
            for exp in resume_data.get("experience", [])
        ]),
        " ".join([
            proj.get("description", "") + " " + proj.get("title", "")
            for proj in resume_data.get("projects", [])
        ]),
    ]).lower()

    matched = sum(1 for kw in keywords if kw.lower() in text)
    base_score = int((matched / len(keywords)) * 70) if keywords else 0

    # Bonus points
    bonus = 0
    if resume_data.get("summary"): bonus += 5
    if len(resume_data.get("skills", [])) >= 5: bonus += 5
    if len(resume_data.get("experience", [])) >= 1: bonus += 10
    if resume_data.get("education"): bonus += 5
    if resume_data.get("linkedin") or resume_data.get("github"): bonus += 5

    score = min(base_score + bonus, 100)
    matched_kws = [kw for kw in keywords if kw.lower() in text]
    missing_kws = [kw for kw in keywords if kw.lower() not in text][:10]

    return {
        "score": score,
        "domain": domain,
        "matched_keywords": matched_kws,
        "missing_keywords": missing_kws,
        "breakdown": {
            "keyword_score": base_score,
            "bonus_score": bonus,
        }
    }


def compute_template_scores(resume_data):
    """Score resume against each company template style."""
    templates = {
        "Google":    {"weight_skills": 0.4, "weight_exp": 0.4, "weight_edu": 0.2},
        "Microsoft": {"weight_skills": 0.3, "weight_exp": 0.5, "weight_edu": 0.2},
        "Meta":      {"weight_skills": 0.5, "weight_exp": 0.35, "weight_edu": 0.15},
        "Oracle":    {"weight_skills": 0.35, "weight_exp": 0.45, "weight_edu": 0.2},
    }
    scores = {}
    skills_count = min(len(resume_data.get("skills", [])), 15) / 15
    exp_count    = min(len(resume_data.get("experience", [])), 5) / 5
    edu_count    = min(len(resume_data.get("education", [])), 3) / 3

    for tpl, w in templates.items():
        raw = (w["weight_skills"] * skills_count +
               w["weight_exp"]   * exp_count +
               w["weight_edu"]   * edu_count)
        scores[tpl] = int(raw * 100)
    return scores


# ══════════════════════════════════════════════════════════════════════════════
#  AUTH ROUTES
# ══════════════════════════════════════════════════════════════════════════════

@app.route("/")
def index():
    if "user_id" in session:
        return redirect(url_for("dashboard"))
    return render_template("landing.html")


@app.route("/register", methods=["GET", "POST"])
def register():
    if "user_id" in session:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        name  = request.form.get("name", "").strip()
        email = request.form.get("email", "").strip().lower()
        pwd   = request.form.get("password", "")
        if not name or not email or not pwd:
            flash("All fields are required.", "error")
            return render_template("register.html")
        if len(pwd) < 6:
            flash("Password must be at least 6 characters.", "error")
            return render_template("register.html")
        if users_col.find_one({"email": email}):
            flash("Email already registered.", "error")
            return render_template("register.html")
        hashed = bcrypt.hashpw(pwd.encode(), bcrypt.gensalt())
        user = {
            "name": name,
            "email": email,
            "password": hashed,
            "auth_provider": "email",
            "avatar": None,
            "created_at": datetime.utcnow(),
        }
        result = users_col.insert_one(user)
        session["user_id"]    = str(result.inserted_id)
        session["user_name"]  = name
        session["user_email"] = email
        flash("Account created successfully!", "success")
        return redirect(url_for("dashboard"))
    return render_template("register.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if "user_id" in session:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        pwd   = request.form.get("password", "")
        user  = users_col.find_one({"email": email})
        if not user or user.get("auth_provider") != "email":
            flash("Invalid credentials.", "error")
            return render_template("login.html")
        if not bcrypt.checkpw(pwd.encode(), user["password"]):
            flash("Invalid credentials.", "error")
            return render_template("login.html")
        session["user_id"]    = str(user["_id"])
        session["user_name"]  = user.get("name", "")
        session["user_email"] = user.get("email", "")
        return redirect(url_for("dashboard"))
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("index"))


# ── Google OAuth ─────────────────────────────────────────────────────────────
@app.route("/auth/google")
def google_login():
    oauth = OAuth2Session(GOOGLE_CLIENT_ID, redirect_uri=GOOGLE_REDIRECT_URI,
                          scope=["openid", "email", "profile"])
    uri, state = oauth.create_authorization_url(GOOGLE_AUTH_URL,
                                                access_type="offline",
                                                prompt="select_account")
    session["oauth_state"] = state
    return redirect(uri)


@app.route("/auth/google/callback")
def google_callback():
    state = session.pop("oauth_state", None)
    if not state:
        flash("OAuth state missing.", "error")
        return redirect(url_for("login"))
    oauth = OAuth2Session(GOOGLE_CLIENT_ID, redirect_uri=GOOGLE_REDIRECT_URI,
                          state=state)
    try:
        token = oauth.fetch_token(
            GOOGLE_TOKEN_URL,
            client_secret=GOOGLE_CLIENT_SECRET,
            authorization_response=request.url,
        )
        userinfo = oauth.get(GOOGLE_USERINFO_URL).json()
    except Exception as e:
        app.logger.error(f"Google OAuth error: {e}")
        flash("Google login failed.", "error")
        return redirect(url_for("login"))

    email  = userinfo.get("email", "").lower()
    name   = userinfo.get("name", email)
    avatar = userinfo.get("picture", "")

    user = users_col.find_one({"email": email})
    if user:
        users_col.update_one({"_id": user["_id"]},
                             {"$set": {"auth_provider": "google", "avatar": avatar}})
    else:
        result = users_col.insert_one({
            "name": name, "email": email,
            "auth_provider": "google", "avatar": avatar,
            "created_at": datetime.utcnow(),
        })
        user = users_col.find_one({"_id": result.inserted_id})

    session["user_id"]    = str(user["_id"])
    session["user_name"]  = name
    session["user_email"] = email
    return redirect(url_for("dashboard"))


# ══════════════════════════════════════════════════════════════════════════════
#  DASHBOARD
# ══════════════════════════════════════════════════════════════════════════════

@app.route("/dashboard")
@login_required
def dashboard():
    uid = session["user_id"]
    resumes = list(resumes_col.find({"user_id": uid}).sort("updated_at", -1))
    for r in resumes:
        r["_id"] = str(r["_id"])
    user = get_current_user()
    return render_template("dashboard.html", resumes=resumes, user=user)


# ══════════════════════════════════════════════════════════════════════════════
#  RESUME CRUD
# ══════════════════════════════════════════════════════════════════════════════

@app.route("/resume/new")
@login_required
def new_resume():
    return render_template("builder.html", resume=None, resume_id=None)


@app.route("/resume/<resume_id>/edit")
@login_required
def edit_resume(resume_id):
    try:
        resume = resumes_col.find_one({
            "_id": ObjectId(resume_id),
            "user_id": session["user_id"]
        })
    except Exception:
        flash("Resume not found.", "error")
        return redirect(url_for("dashboard"))
    if not resume:
        flash("Resume not found.", "error")
        return redirect(url_for("dashboard"))
    resume["_id"] = str(resume["_id"])
    return render_template("builder.html", resume=resume, resume_id=resume_id)


@app.route("/api/resume", methods=["POST"])
@login_required
def api_create_resume():
    data = request.get_json(silent=True) or {}
    now  = datetime.utcnow()
    resume = {
        "user_id":        session["user_id"],
        "title":          data.get("title", "Untitled Resume"),
        "template":       data.get("template", "google"),
        "full_name":      data.get("full_name", ""),
        "email":          data.get("email", ""),
        "phone":          data.get("phone", ""),
        "location":       data.get("location", ""),
        "linkedin":       data.get("linkedin", ""),
        "github":         data.get("github", ""),
        "website":        data.get("website", ""),
        "summary":        data.get("summary", ""),
        "skills":         data.get("skills", []),
        "experience":     data.get("experience", []),
        "education":      data.get("education", []),
        "projects":       data.get("projects", []),
        "certifications": data.get("certifications", []),
        "ats_score":      data.get("ats_score", None),
        "ats_domain":     data.get("ats_domain", "Software Engineering"),
        "created_at":     now,
        "updated_at":     now,
    }
    result = resumes_col.insert_one(resume)
    return jsonify({"success": True, "resume_id": str(result.inserted_id)}), 201


@app.route("/api/resume/<resume_id>", methods=["GET"])
@login_required
def api_get_resume(resume_id):
    try:
        resume = resumes_col.find_one({
            "_id": ObjectId(resume_id),
            "user_id": session["user_id"]
        })
    except Exception:
        return jsonify({"error": "Invalid ID"}), 400
    if not resume:
        return jsonify({"error": "Not found"}), 404
    resume["_id"] = str(resume["_id"])
    return jsonify(resume)


@app.route("/api/resume/<resume_id>", methods=["PUT"])
@login_required
def api_update_resume(resume_id):
    data = request.get_json(silent=True) or {}
    try:
        existing = resumes_col.find_one({
            "_id": ObjectId(resume_id),
            "user_id": session["user_id"]
        })
    except Exception:
        return jsonify({"error": "Invalid ID"}), 400
    if not existing:
        return jsonify({"error": "Not found"}), 404

    update = {
        "title":          data.get("title",          existing.get("title", "Untitled")),
        "template":       data.get("template",       existing.get("template", "google")),
        "full_name":      data.get("full_name",      existing.get("full_name", "")),
        "email":          data.get("email",          existing.get("email", "")),
        "phone":          data.get("phone",          existing.get("phone", "")),
        "location":       data.get("location",       existing.get("location", "")),
        "linkedin":       data.get("linkedin",       existing.get("linkedin", "")),
        "github":         data.get("github",         existing.get("github", "")),
        "website":        data.get("website",        existing.get("website", "")),
        "summary":        data.get("summary",        existing.get("summary", "")),
        "skills":         data.get("skills",         existing.get("skills", [])),
        "experience":     data.get("experience",     existing.get("experience", [])),
        "education":      data.get("education",      existing.get("education", [])),
        "projects":       data.get("projects",       existing.get("projects", [])),
        "certifications": data.get("certifications", existing.get("certifications", [])),
        "ats_score":      data.get("ats_score",      existing.get("ats_score")),
        "ats_domain":     data.get("ats_domain",     existing.get("ats_domain", "Software Engineering")),
        "updated_at":     datetime.utcnow(),
    }
    resumes_col.update_one({"_id": ObjectId(resume_id)}, {"$set": update})
    return jsonify({"success": True})


@app.route("/api/resume/<resume_id>", methods=["DELETE"])
@login_required
def api_delete_resume(resume_id):
    try:
        result = resumes_col.delete_one({
            "_id": ObjectId(resume_id),
            "user_id": session["user_id"]
        })
    except Exception:
        return jsonify({"error": "Invalid ID"}), 400
    if result.deleted_count == 0:
        return jsonify({"error": "Not found"}), 404
    return jsonify({"success": True})


@app.route("/api/resumes", methods=["GET"])
@login_required
def api_list_resumes():
    resumes = list(resumes_col.find({"user_id": session["user_id"]}).sort("updated_at", -1))
    for r in resumes:
        r["_id"] = str(r["_id"])
        r.pop("password", None)
    return jsonify(resumes)


# ══════════════════════════════════════════════════════════════════════════════
#  UPLOAD & PARSE
# ══════════════════════════════════════════════════════════════════════════════

@app.route("/upload")
@login_required
def upload_page():
    return render_template("upload.html")


RESUME_SCHEMA = {
    "type": "object",
    "properties": {
        "full_name": {"type": "string"},
        "email": {"type": "string"},
        "phone": {"type": "string"},
        "location": {"type": "string"},
        "linkedin": {"type": "string"},
        "github": {"type": "string"},
        "website": {"type": "string"},
        "summary": {"type": "string"},
        "skills": {"type": "array", "items": {"type": "string"}},
        "experience": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "title": {"type": "string"},
                    "company": {"type": "string"},
                    "duration": {"type": "string"},
                    "location": {"type": "string"},
                    "description": {"type": "string"}
                },
                "required": ["title", "company", "duration", "location", "description"],
                "additionalProperties": False
            }
        },
        "education": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "degree": {"type": "string"},
                    "institution": {"type": "string"},
                    "year": {"type": "string"},
                    "gpa": {"type": "string"}
                },
                "required": ["degree", "institution", "year", "gpa"],
                "additionalProperties": False
            }
        },
        "projects": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "title": {"type": "string"},
                    "technologies": {"type": "string"},
                    "description": {"type": "string"},
                    "link": {"type": "string"}
                },
                "required": ["title", "technologies", "description", "link"],
                "additionalProperties": False
            }
        },
        "certifications": {"type": "array", "items": {"type": "string"}}
    },
    "required": [
        "full_name", "email", "phone", "location", "linkedin", "github",
        "website", "summary", "skills", "experience", "education",
        "projects", "certifications"
    ],
    "additionalProperties": False
}

RESUME_RESPONSE_FORMAT = {
    "type": "json_schema",
    "json_schema": {
        "name": "resume_parser",
        "strict": True,
        "schema": RESUME_SCHEMA
    }
}


def clean_extracted_text(text):
    """Normalize extracted resume text while preserving useful section structure."""
    if not text:
        return ""
    text = text.replace("\x00", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text_from_pdf(path):
    """Extract all usable PDF text, with OCR fallback for scanned/image PDFs."""
    page_texts = []

    # PyMuPDF gives better page/block ordering than the legacy parsers.
    if PYMUPDF_AVAILABLE:
        try:
            doc = fitz.open(path)
            for page in doc:
                blocks = page.get_text("blocks")
                if blocks:
                    ordered = sorted(blocks, key=lambda b: (round(b[1], 1), round(b[0], 1)))
                    page_text = "\n".join(
                        str(block[4]).strip() for block in ordered
                        if len(block) >= 5 and str(block[4]).strip()
                    )
                else:
                    page_text = page.get_text("text") or ""
                page_texts.append(page_text)
            doc.close()
        except Exception as e:
            app.logger.warning(f"PyMuPDF extraction failed: {e}")

    text = clean_extracted_text("\n\n".join(page_texts))
    compact_len = len(re.sub(r"\s+", "", text))

    # Scanned PDFs often have almost no embedded text. OCR only when needed.
    if compact_len < 120 and OCR_AVAILABLE and shutil.which("tesseract"):
        try:
            doc = fitz.open(path) if PYMUPDF_AVAILABLE else None
            if doc is not None:
                ocr_pages = []
                for page in doc:
                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    ocr_pages.append(pytesseract.image_to_string(img))
                doc.close()
                ocr_text = clean_extracted_text("\n\n".join(ocr_pages))
                if len(re.sub(r"\s+", "", ocr_text)) > compact_len:
                    text = ocr_text
        except Exception as e:
            app.logger.warning(f"PDF OCR failed: {e}")

    # Legacy fallbacks.
    if len(re.sub(r"\s+", "", text)) < 120 and PDFMINER_AVAILABLE:
        try:
            fallback = clean_extracted_text(pdfminer_extract(path) or "")
            if len(re.sub(r"\s+", "", fallback)) > len(re.sub(r"\s+", "", text)):
                text = fallback
        except Exception as e:
            app.logger.warning(f"pdfminer extraction failed: {e}")

    if len(re.sub(r"\s+", "", text)) < 120 and PDF_AVAILABLE:
        try:
            with open(path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                fallback = clean_extracted_text(
                    "\n".join((page.extract_text() or "") for page in reader.pages)
                )
            if len(re.sub(r"\s+", "", fallback)) > len(re.sub(r"\s+", "", text)):
                text = fallback
        except Exception as e:
            app.logger.warning(f"PyPDF2 extraction failed: {e}")

    return text


def extract_text_from_docx(path):
    """Extract DOCX paragraphs plus table content."""
    if not DOCX_AVAILABLE:
        return ""

    try:
        doc = DocxDocument(path)
        parts = []

        for p in doc.paragraphs:
            value = p.text.strip()
            if value:
                parts.append(value)

        for table in doc.tables:
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    value = " ".join(cell.text.split()).strip()
                    if value:
                        cells.append(value)
                if cells:
                    parts.append(" | ".join(cells))

        return clean_extracted_text("\n".join(parts))
    except Exception as e:
        app.logger.error(f"DOCX extraction failed: {e}")
        return ""


def _to_text(value):
    """Coerce model values to the string types expected by the builder."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, list):
        return "\n".join(str(item).strip() for item in value if str(item).strip())
    return str(value).strip()


def normalize_resume_data(data):
    """Normalize structured AI output before returning it to the frontend."""
    if not isinstance(data, dict):
        return {
            "full_name": "", "email": "", "phone": "", "location": "",
            "linkedin": "", "github": "", "website": "", "summary": "",
            "skills": [], "experience": [], "education": [],
            "projects": [], "certifications": []
        }

    normalized = {
        "full_name": _to_text(data.get("full_name")),
        "email": _to_text(data.get("email")),
        "phone": _to_text(data.get("phone")),
        "location": _to_text(data.get("location")),
        "linkedin": _to_text(data.get("linkedin")),
        "github": _to_text(data.get("github")),
        "website": _to_text(data.get("website")),
        "summary": _to_text(data.get("summary")),
        "skills": [],
        "experience": [],
        "education": [],
        "projects": [],
        "certifications": [],
    }

    skills = data.get("skills", [])
    if isinstance(skills, list):
        normalized["skills"] = list(dict.fromkeys(
            _to_text(item) for item in skills if _to_text(item)
        ))

    for item in data.get("experience", []) if isinstance(data.get("experience"), list) else []:
        if not isinstance(item, dict):
            continue
        normalized["experience"].append({
            "title": _to_text(item.get("title")),
            "company": _to_text(item.get("company")),
            "duration": _to_text(item.get("duration")),
            "location": _to_text(item.get("location")),
            "description": _to_text(item.get("description")),
        })

    for item in data.get("education", []) if isinstance(data.get("education"), list) else []:
        if not isinstance(item, dict):
            continue
        normalized["education"].append({
            "degree": _to_text(item.get("degree")),
            "institution": _to_text(item.get("institution")),
            "year": _to_text(item.get("year")),
            "gpa": _to_text(item.get("gpa")),
        })

    for item in data.get("projects", []) if isinstance(data.get("projects"), list) else []:
        if not isinstance(item, dict):
            continue
        normalized["projects"].append({
            "title": _to_text(item.get("title")),
            "technologies": _to_text(item.get("technologies")),
            "description": _to_text(item.get("description")),
            "link": _to_text(item.get("link")),
        })

    certifications = data.get("certifications", [])
    if isinstance(certifications, list):
        normalized["certifications"] = list(dict.fromkeys(
            _to_text(item) for item in certifications if _to_text(item)
        ))

    return normalized


@app.route("/api/upload", methods=["POST"])
@login_required
def api_upload_resume():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "Empty filename"}), 400

    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in [".pdf", ".docx"]:
        return jsonify({"error": "Only PDF and DOCX files are supported"}), 400

    # Keep uploads isolated per request to avoid concurrent-user collisions.
    suffix = ext
    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix, prefix="resume_upload_") as tmp:
            tmp_path = tmp.name
            file.save(tmp_path)

        if ext == ".pdf":
            raw_text = extract_text_from_pdf(tmp_path)
        else:
            raw_text = extract_text_from_docx(tmp_path)
    finally:
        if tmp_path:
            try:
                os.remove(tmp_path)
            except OSError:
                pass

    raw_text = clean_extracted_text(raw_text)
    if not raw_text:
        return jsonify({
            "error": (
                "Could not extract readable text from this file. "
                "For scanned/image-only PDFs, enable Tesseract OCR on the server."
            )
        }), 422

    # GPT-OSS 120B supports a 131K context window. Keep a generous safety cap
    # so unusually large uploads cannot consume the entire request budget.
    resume_text = raw_text[:60000]

    prompt = f"""Extract the candidate information from the resume text below.

Rules:
- Use ONLY information explicitly present in the resume.
- Never invent employers, dates, degrees, skills, links, metrics, or credentials.
- Preserve all relevant work experiences, education entries, projects, and certifications found in the text.
- Keep each experience/project as a separate array item.
- If a field is missing, return an empty string or empty array.
- Return every required field exactly as defined by the schema.

Resume text:
{resume_text}"""

    parsed_text = groq_chat(
        [
            {
                "role": "system",
                "content": (
                    "You are a high-precision resume information extraction system. "
                    "Extract structured facts faithfully. Do not summarize away important entries."
                ),
            },
            {"role": "user", "content": prompt},
        ],
        temperature=0,
        max_tokens=5000,
        response_format=RESUME_RESPONSE_FORMAT,
    )

    if not parsed_text:
        return jsonify({
            "error": "AI parsing failed. Please try again or check the GROQ_MODEL/GROQ_API_KEY configuration."
        }), 502

    try:
        parsed = normalize_resume_data(json.loads(parsed_text))
    except Exception as e:
        app.logger.error(f"Resume JSON parsing failed: {e}")
        return jsonify({
            "error": "The AI returned an invalid resume structure. Please try again."
        }), 502

    return jsonify({"success": True, "data": parsed}), 200


# ══════════════════════════════════════════════════════════════════════════════
#  ATS SCORING
# ══════════════════════════════════════════════════════════════════════════════

@app.route("/api/ats-score", methods=["POST"])
@login_required
def api_ats_score():
    data   = request.get_json(silent=True) or {}
    domain = data.get("domain", "Software Engineering")
    result = compute_ats_score(data, domain)
    result["template_scores"] = compute_template_scores(data)

    resume_id = data.get("resume_id")
    if resume_id:
        try:
            resumes_col.update_one(
                {"_id": ObjectId(resume_id), "user_id": session["user_id"]},
                {"$set": {
                    "ats_score":  result["score"],
                    "ats_domain": domain,
                    "ats_details": result,
                    "updated_at": datetime.utcnow()
                }}
            )
        except Exception:
            pass

    return jsonify(result)


# ══════════════════════════════════════════════════════════════════════════════
#  AI SUGGESTIONS
# ══════════════════════════════════════════════════════════════════════════════

@app.route("/api/ai/suggestions", methods=["POST"])
@login_required
def api_ai_suggestions():
    data   = request.get_json(silent=True) or {}
    domain = data.get("domain", "Software Engineering")
    skills = data.get("skills", [])

    ats     = compute_ats_score(data, domain)
    missing = ats.get("missing_keywords", [])

    prompt = f"""You are a professional resume coach for the {domain} domain.

Given this resume summary:
{data.get('summary', 'N/A')[:500]}

Skills already present: {', '.join(skills[:20]) if skills else 'None'}
Missing ATS keywords: {', '.join(missing)}

Provide exactly 3 concise, actionable improvement suggestions (one sentence each) and list 8 specific skills to add.
Return ONLY valid JSON:
{{
  "suggestions": ["suggestion1", "suggestion2", "suggestion3"],
  "skills_to_add": ["skill1", "skill2", "skill3", "skill4", "skill5", "skill6", "skill7", "skill8"]
}}"""

    result_text = groq_chat([{"role": "user", "content": prompt}], temperature=0.6)

    try:
        cleaned = result_text.strip()
        if cleaned.startswith("```"):
            cleaned = cleaned.split("```")[1]
            if cleaned.startswith("json"):
                cleaned = cleaned[4:]
        result = json.loads(cleaned.strip())
    except Exception:
        result = {
            "suggestions": [
                "Quantify your achievements with specific metrics and numbers.",
                "Add more industry-specific keywords to improve ATS matching.",
                "Expand your skills section with trending technologies in your domain."
            ],
            "skills_to_add": missing[:8] if missing else []
        }

    result["missing_keywords"] = missing
    return jsonify(result)


# ── AI section rewrite ────────────────────────────────────────────────────────
@app.route("/api/ai/rewrite", methods=["POST"])
@login_required
def api_ai_rewrite():
    data    = request.get_json(silent=True) or {}
    section = data.get("section", "summary")
    content = data.get("content", "")
    context = data.get("context", "")

    section_prompts = {
        "summary": f"""Rewrite this professional summary to be more impactful, concise, and ATS-friendly (3-4 sentences max):

Current summary: {content}
Context about the person: {context}

Return ONLY the improved summary text, no explanations.""",

        "experience": f"""Rewrite this work experience description using strong action verbs, quantified achievements, and ATS keywords:

Current description: {content}
Job context: {context}

Return ONLY bullet points (start each with •), no other text.""",

        "skills": f"""Based on the context, suggest an optimized skills list in comma-separated format:

Current skills: {content}
Role/domain context: {context}

Return ONLY a comma-separated list of skills, nothing else.""",

        "projects": f"""Rewrite this project description to highlight technical impact and skills:

Current description: {content}
Technologies used: {context}

Return ONLY the improved description, no explanations.""",

        "education": f"""Improve this education entry to highlight relevant coursework and achievements:

Current text: {content}
Context: {context}

Return ONLY the improved text, no explanations.""",
    }

    prompt = section_prompts.get(section, section_prompts["summary"])
    result = groq_chat([{"role": "user", "content": prompt}], temperature=0.7)

    if not result:
        result = content  # fallback to original

    return jsonify({"rewritten": result})


# ══════════════════════════════════════════════════════════════════════════════
#  STATIC TEMPLATE PREVIEW
# ══════════════════════════════════════════════════════════════════════════════

@app.route("/resume/<resume_id>/preview")
@login_required
def preview_resume(resume_id):
    try:
        resume = resumes_col.find_one({
            "_id": ObjectId(resume_id),
            "user_id": session["user_id"]
        })
    except Exception:
        return redirect(url_for("dashboard"))
    if not resume:
        return redirect(url_for("dashboard"))
    resume["_id"] = str(resume["_id"])
    return render_template("preview.html", resume=resume)


# ══════════════════════════════════════════════════════════════════════════════
#  ERROR HANDLERS
# ══════════════════════════════════════════════════════════════════════════════

@app.errorhandler(404)
def not_found(e):
    return render_template("error.html", code=404, message="Page not found"), 404


@app.errorhandler(500)
def server_error(e):
    return render_template("error.html", code=500, message="Internal server error"), 500


# ══════════════════════════════════════════════════════════════════════════════
#  MAIN
# ══════════════════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    app.run(debug=True, port=5000)
